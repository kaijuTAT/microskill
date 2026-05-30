from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import pdfplumber
from docx import Document
from pptx import Presentation


VIDEO_RE = re.compile(r"-P(\d+)\s+(.+?)\.ai-zh\.srt$", re.IGNORECASE)


def read_metadata(video_root: Path) -> list[dict]:
    metadata_files = sorted(video_root.rglob("*元数据.json"))
    if not metadata_files:
        return []
    with metadata_files[0].open("r", encoding="utf-8") as f:
        data = json.load(f)
    pages = data.get("pages", [])
    return [
        {
            "part": page.get("page"),
            "title": page.get("part", ""),
            "duration": int(page.get("duration", 0)),
        }
        for page in pages
    ]


def count_subtitle_lines(path: Path) -> int:
    count = 0
    with path.open("r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            line = line.strip()
            if line and not line.isdigit() and "-->" not in line:
                count += 1
    return count


def scan_subtitles(video_root: Path) -> list[dict]:
    rows = []
    for path in sorted(video_root.rglob("*.srt")):
        match = VIDEO_RE.search(path.name)
        part = int(match.group(1)) if match else None
        title = match.group(2) if match else path.stem
        rows.append(
            {
                "part": part,
                "title": title,
                "lines": count_subtitle_lines(path),
                "path": str(path),
            }
        )
    return sorted(rows, key=lambda row: (row["part"] is None, row["part"] or 9999, row["title"]))


def scan_course_files(root: Path) -> list[dict]:
    wanted = {".pptx", ".pdf", ".docx"}
    rows = []
    for path in sorted(root.iterdir()):
        if path.name.startswith("._") or path.suffix.lower() not in wanted:
            continue
        rows.append({"name": path.name, "kind": path.suffix.lower(), "size": path.stat().st_size})
    return rows


def extract_courseware_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        prs = Presentation(str(path))
        parts = []
        for slide_no, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if text and text.strip():
                    texts.append(text.strip())
            if texts:
                parts.append(f"[slide {slide_no}] " + " ".join(texts))
        return "\n".join(parts)
    if suffix == ".docx":
        doc = Document(str(path))
        return "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    if suffix == ".pdf":
        parts = []
        with pdfplumber.open(str(path)) as pdf:
            for page_no, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(f"[page {page_no}] {text.strip()}")
        return "\n".join(parts)
    return ""


def fmt_duration(seconds: int) -> str:
    minutes, sec = divmod(seconds, 60)
    return f"{minutes}m{sec:02d}s"


def write_markdown(root: Path, output: Path) -> None:
    video_root = root / "网课"
    metadata = read_metadata(video_root)
    subtitles = scan_subtitles(video_root)
    files = scan_course_files(root)

    lines = [
        "# Generated Course Index",
        "",
        f"- Workspace: `{root}`",
        f"- Video folder: `{video_root}`",
        f"- Lessons from metadata: {len(metadata)}",
        f"- Subtitle files: {len(subtitles)}",
        f"- Courseware files: {len(files)}",
        "",
        "## Lesson Map",
        "",
        "| Part | Topic | Duration | Subtitle Lines |",
        "|---|---|---:|---:|",
    ]

    subtitle_lines = {row["part"]: row["lines"] for row in subtitles if row["part"] is not None}
    for row in metadata:
        part = row["part"]
        lines.append(
            f"| P{part:02d} | {row['title']} | {fmt_duration(row['duration'])} | {subtitle_lines.get(part, '')} |"
        )

    lines += ["", "## Courseware", "", "| File | Type | Size |", "|---|---|---:|"]
    for row in files:
        lines.append(f"| {row['name']} | {row['kind']} | {row['size']} |")

    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


FOCUS_TOPICS = {
    "assembly": {
        "parts": set(range(17, 25)),
        "keywords": ["汇编", "伪指令", "程序", "数据段", "代码段", "DOS", "循环", "转移"],
    },
    "8259": {
        "parts": {29, 30, 31},
        "keywords": ["8259", "中断", "ICW", "OCW", "IRR", "IMR", "ISR", "优先级"],
    },
    "8255": {
        "parts": {32, 33, 34, 35, 36},
        "keywords": ["8255", "并行", "A口", "B口", "C口", "控制字", "方式", "握手"],
    },
    "8253": {
        "parts": {38, 39, 40},
        "keywords": ["8253", "定时", "计数", "控制字", "方式", "初值", "频率", "时钟"],
    },
}


def read_srt_text(path: Path) -> list[str]:
    rows = []
    with path.open("r", encoding="utf-8", errors="ignore") as f:
        current_time = ""
        for line in f:
            line = line.strip()
            if not line or line.isdigit():
                continue
            if "-->" in line:
                current_time = line.split("-->", 1)[0].strip()
                continue
            rows.append(f"{current_time} {line}".strip())
    return rows


def build_focus_evidence(root: Path, output: Path) -> None:
    video_root = root / "网课"
    metadata = {row["part"]: row for row in read_metadata(video_root)}
    subtitles = scan_subtitles(video_root)
    srt_by_part = {row["part"]: Path(row["path"]) for row in subtitles if row["part"] is not None}

    lines = [
        "# Focus Evidence Index",
        "",
        "This file points the skill to the highest-value source lessons for the user's exam scope.",
        "",
    ]

    for topic, config in FOCUS_TOPICS.items():
        lines += [f"## {topic}", ""]
        for part in sorted(config["parts"]):
            meta = metadata.get(part, {})
            srt_path = srt_by_part.get(part)
            title = meta.get("title", f"P{part:02d}")
            duration = fmt_duration(meta.get("duration", 0)) if meta else ""
            lines.append(f"### P{part:02d} {title} ({duration})")
            if srt_path:
                lines.append(f"- Subtitle: `{srt_path}`")
                hits = []
                for row in read_srt_text(srt_path):
                    if any(keyword.lower() in row.lower() for keyword in config["keywords"]):
                        hits.append(row)
                    if len(hits) >= 8:
                        break
                if hits:
                    lines.append("- Keyword hits:")
                    for hit in hits:
                        lines.append(f"  - {hit}")
                else:
                    lines.append("- Keyword hits: none in first pass")
            lines.append("")

    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_courseware_evidence(root: Path, output: Path) -> None:
    files = [root / row["name"] for row in scan_course_files(root)]
    lines = [
        "# Courseware Evidence Index",
        "",
        "This file lists courseware and practice-file snippets related to the user's priority topics.",
        "",
    ]
    for topic, config in FOCUS_TOPICS.items():
        lines += [f"## {topic}", ""]
        found_any = False
        for path in files:
            try:
                text = extract_courseware_text(path)
            except Exception as exc:
                lines.append(f"### {path.name}")
                lines.append(f"- Extraction error: {exc}")
                lines.append("")
                continue
            snippets = []
            for row in text.splitlines():
                if any(keyword.lower() in row.lower() for keyword in config["keywords"]):
                    snippets.append(row[:260])
                if len(snippets) >= 5:
                    break
            if snippets:
                found_any = True
                lines.append(f"### {path.name}")
                for snippet in snippets:
                    lines.append(f"- {snippet}")
                lines.append("")
        if not found_any:
            lines.append("- No direct keyword snippets found.")
            lines.append("")

    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a compact index for HEU microcomputer review materials.")
    parser.add_argument("--root", default=r"D:\微机原理", help="Course workspace root.")
    parser.add_argument("--output", default="generated-course-index.md", help="Output Markdown path.")
    parser.add_argument("--focus-output", help="Optional focus evidence Markdown path.")
    parser.add_argument("--courseware-output", help="Optional courseware evidence Markdown path.")
    args = parser.parse_args()

    root = Path(args.root)
    output = Path(args.output)
    write_markdown(root, output)
    print(f"Wrote {output}")
    if args.focus_output:
        focus_output = Path(args.focus_output)
        build_focus_evidence(root, focus_output)
        print(f"Wrote {focus_output}")
    if args.courseware_output:
        courseware_output = Path(args.courseware_output)
        build_courseware_evidence(root, courseware_output)
        print(f"Wrote {courseware_output}")


if __name__ == "__main__":
    main()
